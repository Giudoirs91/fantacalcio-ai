import openpyxl
import json
import unicodedata
import re

with open('data/processed/processed_players_master.json', 'r', encoding='utf-8') as f:
    master_players = json.load(f)

wb = openpyxl.load_workbook('data/raw/fantarefri-rosters-1789316126497.xlsx')
sheet = wb['ROSE']
headers = list(sheet.iter_rows(values_only=True))[0]
all_rows = list(sheet.iter_rows(values_only=True))

def clean_str(s):
    if not s: return ''
    s = unicodedata.normalize('NFD', str(s)).encode('ascii', 'ignore').decode('utf-8')
    s = re.sub(r'[^a-zA-Z0-9]', ' ', s.lower())
    return re.sub(r'\s+', ' ', s).strip()

lookup_exact = {}
for p in master_players:
    c = clean_str(p['name'])
    lookup_exact[c] = p

def match_player(raw):
    craw = clean_str(raw)
    if not craw:
        return None
    if craw in lookup_exact:
        return lookup_exact[craw]
    
    tokens = craw.split()
    surname = tokens[0]
    # Exact surname matches in master
    candidates = [p for p in master_players if clean_str(p['name']).startswith(surname + ' ') or clean_str(p['name']) == surname]
    if len(candidates) == 1:
        return candidates[0]
    elif len(candidates) > 1:
        # Check second token as initial or second surname
        if len(tokens) > 1:
            second = tokens[1]
            for c in candidates:
                c_parts = clean_str(c['name']).split()
                if len(c_parts) > 1 and c_parts[1].startswith(second[0]):
                    return c
        return candidates[0]
    
    # Substring search if not found
    for p in master_players:
        cp = clean_str(p['name'])
        if len(craw) >= 4 and (craw in cp or cp.startswith(craw)):
            return p
            
    # Fuzzy last resort
    for p in master_players:
        cp_parts = clean_str(p['name']).split()
        if cp_parts and cp_parts[0] in craw:
            return p
            
    return None

total = 0
matched = 0
unmatched = []

for col_idx in range(0, len(headers)):
    h = headers[col_idx]
    if h and str(h).strip() and str(h).strip().lower() not in ['costo', 'prezzo', 'ruolo', 'r', 'none']:
        t_name = str(h).strip()
        for r_idx in range(1, len(all_rows)):
            row = all_rows[r_idx]
            p_name = row[col_idx] if col_idx < len(row) else None
            p_cost = row[col_idx + 1] if col_idx + 1 < len(row) else 1
            if p_name and str(p_name).strip():
                total += 1
                raw = str(p_name).strip()
                res = match_player(raw)
                if res:
                    matched += 1
                else:
                    unmatched.append((t_name, raw))

print(f"Total: {total}, Matched: {matched}/{total} ({matched/total*100:.1f}%)")
if unmatched:
    print("Unmatched:", unmatched)
else:
    print("ALL 100% MATCHED!")
