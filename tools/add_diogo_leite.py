import os
import openpyxl
import pandas as pd
import json

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def add_diogo_leite():
    excel_path = os.path.join(ROOT_DIR, "data", "raw", "Listone_definitivo_Fantacalcio_Stagione_2026_27.xlsx")
    print(f"1. Aggiunta a {excel_path}...")
    wb = openpyxl.load_workbook(excel_path)
    
    player_id = 7629
    row_data = [
        player_id,      # Id
        'D',            # R
        'Dc',           # RM (Mantra)
        'Diogo Leite',  # Nome
        'Lazio',        # Squadra
        11,             # Qt.A
        11,             # Qt.I
        0,              # Diff.
        11,             # Qt.A M
        11,             # Qt.I M
        0,              # Diff.M
        16,             # FVM
        16              # FVM M
    ]
    
    # Aggiungi a 'Tutti' se non presente
    sheet_tutti = wb['Tutti']
    existing = False
    for r in range(2, sheet_tutti.max_row + 1):
        if sheet_tutti.cell(r, 4).value == 'Diogo Leite':
            existing = True
            break
    
    if not existing:
        sheet_tutti.append(row_data)
        sheet_dif = wb['Difensori']
        sheet_dif.append(row_data)
        wb.save(excel_path)
        print("   -> Salvato con successo nel Listone Excel!")
    else:
        print("   -> Diogo Leite già presente nel Listone Excel.")

    # 2. Aggiunta a preview_all_518_players_stats_2025_26.csv
    csv_path = os.path.join(ROOT_DIR, "data", "raw", "preview_all_518_players_stats_2025_26.csv")
    if os.path.exists(csv_path):
        print(f"2. Aggiunta a {csv_path}...")
        df = pd.read_csv(csv_path)
        if not (df['name'] == 'Diogo Leite').any():
            new_row = {
                'id': player_id,
                'name': 'Diogo Leite',
                'team_2627': 'Lazio',
                'role': 'D',
                'mantra': 'Dc',
                'fvm': 16.0,
                'has_data_2526': True,
                'league_2526': 'Bundesliga',
                'rating_2526': 7.08,
                'mins_2526': 2720.0,
                'goals_2526': 1,
                'assists_2526': 1,
                'xg_2526': 1.2,
                'xg90_2526': 0.04,
                'xgot_2526': 1.1,
                'xa90_2526': 0.03,
                'shots90_2526': 0.45,
                'tkl_int90_2526': 3.8,
                'goals_prevented_2526': None
            }
            df = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)
            df.to_csv(csv_path, index=False)
            print("   -> Salvato con successo in preview_all_518_players_stats_2025_26.csv!")
        else:
            print("   -> Diogo Leite già presente nel CSV statistiche.")

    # 3. Aggiorna config/tactical_db.json e data/raw/database_tattico/tactical_db.json
    for t_path in [
        os.path.join(ROOT_DIR, "config", "tactical_db.json"),
        os.path.join(ROOT_DIR, "data", "raw", "database_tattico", "tactical_db.json")
    ]:
        if os.path.exists(t_path):
            print(f"3. Aggiornamento {t_path}...")
            with open(t_path, "r", encoding="utf-8") as f:
                db = json.load(f)
            
            lazio = db.get("Lazio", {})
            lineup = lazio.get("lineup", [])
            for p in lineup:
                if p.get("pos") == "DC_S":
                    p["name"] = "Diogo Leite"
                    p["sub_name"] = "Provstgaard"
                    p["sub_role"] = "D"
                    p["status"] = "Titolare (65% vs Provstgaard 35%)"
            
            ballottaggi = lazio.get("ballottaggi", [])
            # rimuovi vecchi ballottaggi Provstgaard / Patric
            ballottaggi = [b for b in ballottaggi if b.get("player") not in ["Provstgaard", "Patric"]]
            ballottaggi.append({
                "player": "Diogo Leite",
                "pct": 65,
                "vs": "Provstgaard (35%)"
            })
            ballottaggi.append({
                "player": "Provstgaard",
                "pct": 35,
                "vs": "Diogo Leite (65%)"
            })
            lazio["ballottaggi"] = ballottaggi

            with open(t_path, "w", encoding="utf-8") as f:
                json.dump(db, f, ensure_ascii=False, indent=2)
            print(f"   -> Salvato {t_path}!")

if __name__ == "__main__":
    add_diogo_leite()
